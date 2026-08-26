"""
Script pra carregar a tabela `mapeamento_categoria_tarjado` do Postgres a
partir do de-para já classificado e revisado pelo time ("de-para-categorias-
tarjados - VALIDADO.xlsx", aba "de-para") - mapeia cada combinação distinta
de (SECAO, NEC 1, NEC 2, NEC 3) de `medicamentos_tarjados` (ver
carregar_tarjados.py) pra uma folha de `categorias`. Mesmo papel de
mapeamento_categoria_cmed/mapeamento_categoria_iqvia, com colunas separadas
por nível (secao, nec1, nec2, nec3) - mesmo formato de
mapeamento_categoria_iqvia (tipo_produto/area_farmacia/sub_cat1..4), já que a
taxonomia de origem também tem vários níveis. Mas aqui a classificação por IA
e a revisão humana já vêm prontas no xlsx (feitas fora deste projeto) - não
existe um mapear_categorias_tarjado.py que chama Claude, só esta carga.

Pra cada linha do xlsx, a categoria final é DEP/CAT/SUBCAT (correção manual
do time) quando preenchido, senão departamento/categoria/subcategoria
(sugestão da IA) - e tipo_produto é inferido de tipo_cadastro_sugerido
(Medicamento/Não Medicamento) quando presente, senão descoberto testando
medicamento e depois não_medicamento contra a árvore oficial (sem ambiguidade
na prática: cada combinação departamento/categoria/subcategoria existe em só
um dos dois ramos).

Uso:
    python carregar_mapeamento_tarjado.py criar-tabela
    python carregar_mapeamento_tarjado.py carregar "de-para-categorias-tarjados - VALIDADO.xlsx"

Requer Postgres rodando (docker-compose up -d) e psycopg2-binary instalado.
"""

import argparse

import openpyxl
import psycopg2.extras

import categorias
from db import conectar

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS mapeamento_categoria_tarjado (
    id             BIGSERIAL PRIMARY KEY,
    secao          TEXT,
    nec1           TEXT,
    nec2           TEXT,
    nec3           TEXT,
    categoria_id   BIGINT REFERENCES categorias(id),
    revisado       BOOLEAN NOT NULL DEFAULT false,
    criado_em      TIMESTAMPTZ NOT NULL DEFAULT now(),
    atualizado_em  TIMESTAMPTZ NOT NULL DEFAULT now()
);
CREATE UNIQUE INDEX IF NOT EXISTS mapeamento_categoria_tarjado_chave ON mapeamento_categoria_tarjado (
    coalesce(secao, ''), coalesce(nec1, ''), coalesce(nec2, ''), coalesce(nec3, '')
);
"""


def criar_tabela():
    with conectar() as conn:
        with conn.cursor() as cur:
            cur.execute(SCHEMA_SQL)
        conn.commit()
    print("Tabela criada/confirmada: mapeamento_categoria_tarjado.")


def _resolver_categoria_id(tipo_sugerido, departamento, categoria, subcategoria):
    """Tenta o tipo sugerido pelo xlsx primeiro, senão testa os dois ramos
    oficiais - ver docstring do módulo."""
    if not (departamento and categoria and subcategoria):
        return None
    if tipo_sugerido:
        categoria_id = categorias.resolver_id(tipo_sugerido, departamento, categoria, subcategoria)
        if categoria_id:
            return categoria_id
    for tipo in ("medicamento", "nao_medicamento"):
        if tipo == tipo_sugerido:
            continue
        categoria_id = categorias.resolver_id(tipo, departamento, categoria, subcategoria)
        if categoria_id:
            return categoria_id
    return None


def _tipo_sugerido(valor):
    if valor == "Medicamento":
        return "medicamento"
    if valor and "Não Medicamento" in valor:
        return "nao_medicamento"
    return None


def _linha_para_registro(ws, r):
    secao, nec1, nec2, nec3 = (ws.cell(row=r, column=c).value for c in (1, 2, 3, 4))
    dep_ia, cat_ia, sub_ia = (ws.cell(row=r, column=c).value for c in (6, 7, 8))
    tipo_sugerido_bruto = ws.cell(row=r, column=10).value
    revisado_humanamente = ws.cell(row=r, column=11).value
    dep_h, cat_h, sub_h = (ws.cell(row=r, column=c).value for c in (12, 13, 14))

    departamento = dep_h or dep_ia
    categoria_nome = cat_h or cat_ia
    subcategoria = sub_h or sub_ia
    tipo_sugerido = _tipo_sugerido(tipo_sugerido_bruto)

    categoria_id = _resolver_categoria_id(tipo_sugerido, departamento, categoria_nome, subcategoria)
    revisado = revisado_humanamente == "Sim"
    return secao, nec1, nec2, nec3, categoria_id, revisado


def carregar_mapeamento(caminho_xlsx):
    """
    Lê a aba "de-para" do xlsx e faz upsert na tabela mapeamento_categoria_tarjado
    pela chave natural (secao, nec1, nec2, nec3) - idempotente, pode rodar de
    novo com uma versão mais recente do de-para sem duplicar.
    """
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb["de-para"]

    registros = [_linha_para_registro(ws, r) for r in range(2, ws.max_row + 1)]
    sem_categoria = sum(1 for r in registros if r[4] is None)

    with conectar() as conn:
        with conn.cursor() as cur:
            psycopg2.extras.execute_values(
                cur,
                """
                INSERT INTO mapeamento_categoria_tarjado (secao, nec1, nec2, nec3, categoria_id, revisado)
                VALUES %s
                ON CONFLICT (coalesce(secao, ''), coalesce(nec1, ''), coalesce(nec2, ''), coalesce(nec3, ''))
                DO UPDATE SET categoria_id = EXCLUDED.categoria_id,
                              revisado = EXCLUDED.revisado,
                              atualizado_em = now()
                """,
                registros,
            )
        conn.commit()

    print(f"{len(registros)} combinação(ões) gravada(s) de {caminho_xlsx!r}.")
    if sem_categoria:
        print(f"{sem_categoria} linha(s) sem categoria_id (departamento/categoria/subcategoria vazio ou sem folha correspondente).")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="comando", required=True)

    sub.add_parser("criar-tabela", help="Cria a tabela mapeamento_categoria_tarjado se nao existir")

    p_carregar = sub.add_parser(
        "carregar",
        help='Carrega o de-para validado (aba "de-para") pra mapeamento_categoria_tarjado',
    )
    p_carregar.add_argument(
        "arquivo", help='Caminho do xlsx (ex: "de-para-categorias-tarjados - VALIDADO.xlsx")'
    )

    args = parser.parse_args()

    if args.comando == "criar-tabela":
        criar_tabela()
    elif args.comando == "carregar":
        carregar_mapeamento(args.arquivo)


if __name__ == "__main__":
    main()
